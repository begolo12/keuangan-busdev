import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from backend.database import get_summary, get_members, get_transactions, get_settings, get_connection

def create_styled_cell(ws, row, col, value, font=None, fill=None, alignment=None, border=None, number_format=None):
    cell = ws.cell(row=row, column=col, value=value)
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    if alignment:
        cell.alignment = alignment
    if border:
        cell.border = border
    if number_format:
        cell.number_format = number_format
    return cell

def export_database_to_excel():
    summary = get_summary()
    members = summary["members"]
    settings = get_settings()
    plafon = summary["plafon_induk"]
    tahun = summary["tahun"]

    wb = openpyxl.Workbook()
    # Default sheet
    ws_res = wb.active
    ws_res.title = "Resume"

    # Styles
    thin_border_side = Side(style='thin', color='D0D5DD')
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    thick_bottom = Border(bottom=Side(style='medium', color='1E293B'))

    header_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid') # Deep Navy
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')

    sub_header_fill = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
    sub_header_font = Font(name='Calibri', size=10, bold=True, color='1E293B')

    card_fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
    card_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    title_font = Font(name='Calibri', size=16, bold=True, color='0F172A')
    subtitle_font = Font(name='Calibri', size=11, bold=False, color='475569')
    bold_font = Font(name='Calibri', size=10, bold=True, color='0F172A')
    normal_font = Font(name='Calibri', size=10, color='1E293B')

    rupiah_format = '#,##0'

    # --- RESUME SHEET ---
    ws_res['A1'] = f"TAHUN {tahun}"
    ws_res['A1'].font = Font(name='Calibri', size=11, bold=True, color='64748B')

    ws_res['A3'] = "DASHBOARD REKAPITULASI DANA PERSEKOT, REIMBURSEMENT & OPNAME KAS FISIK"
    ws_res['A3'].font = title_font

    ws_res['A4'] = f"Monitoring Pengeluaran, Distribusi Kas Bon Tim, dan Audit Keseimbangan Fisik vs Catatan (Plafon Kas: Rp {plafon:,.0f})".replace(",", ".")
    ws_res['A4'].font = subtitle_font

    # Kontrol Kas Induk
    ws_res['B6'] = "KONTROL KAS INDUK PERSEKOT"
    ws_res['B6'].font = Font(name='Calibri', size=11, bold=True, color='1E3A8A')

    ws_res['B7'] = "Total Kas Induk Persekot (Plafon Awal)"
    ws_res['D7'] = plafon
    ws_res['D7'].number_format = rupiah_format
    ws_res['D7'].font = bold_font

    ws_res['B8'] = "Total Persekot Didistribusikan ke Tim"
    ws_res['D8'] = f"=D{14 + len(members)}"
    ws_res['D8'].number_format = rupiah_format
    ws_res['D8'].font = bold_font

    ws_res['B9'] = "Sisa Dana di Kas Induk (Belum Dibagi)"
    ws_res['D9'] = "=D7-D8"
    ws_res['D9'].number_format = rupiah_format
    ws_res['D9'].font = bold_font

    ws_res['B10'] = "Status Distribusi Kas Induk"
    ws_res['D10'] = '=IF(D9=0, "100% Teralokasi ke Tim", IF(D9>0, "Sebagian Dana Masih di Kas Induk", "PERINGATAN: Melebihi Plafon Kas!"))'
    ws_res['D10'].font = bold_font

    for r in range(7, 11):
        ws_res.cell(r, 2).font = normal_font
        ws_res.cell(r, 2).fill = card_fill
        ws_res.cell(r, 4).fill = card_fill

    # Rekapitulasi Table Header
    ws_res['A12'] = "REKAPITULASI PENGELUARAN, SALDO BUKU & SISA UANG FISIK RIIL"
    ws_res['A12'].font = Font(name='Calibri', size=11, bold=True, color='1E3A8A')

    headers_rekap = [
        "No.", "Nama Anggota", "Divisi / Pos", "Persekot Diterima (Rp)",
        "Realisasi Belanja (Rp)", "Saldo Catatan Buku (Rp)", "Sisa Uang Fisik Riil (Rp)\n[KOLOM INPUT KAS]",
        "Selisih / Gap Fisik (Rp)\n[Fisik - Buku]", "Status Gap Fisik", "Tindak Lanjut / Action"
    ]

    for c, h in enumerate(headers_rekap, 1):
        cell = ws_res.cell(13, c, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border

    ws_res.row_dimensions[13].height = 30

    row_idx = 14
    member_row_map = {}
    for i, m in enumerate(members, 1):
        member_name = m["name"]
        sheet_safe_name = member_name.replace("'", "")
        member_row_map[member_name] = row_idx

        ws_res.cell(row_idx, 1, i).alignment = Alignment(horizontal='center')
        ws_res.cell(row_idx, 2, member_name).font = bold_font
        ws_res.cell(row_idx, 3, m["divisi"])

        # Formulas linking to member sheet
        ws_res.cell(row_idx, 4, f"='{sheet_safe_name}'!D4").number_format = rupiah_format
        ws_res.cell(row_idx, 5, f"='{sheet_safe_name}'!D5").number_format = rupiah_format
        ws_res.cell(row_idx, 6, f"=D{row_idx}-E{row_idx}").number_format = rupiah_format
        ws_res.cell(row_idx, 7, m["uang_fisik"]).number_format = rupiah_format
        ws_res.cell(row_idx, 8, f"=G{row_idx}-F{row_idx}").number_format = rupiah_format
        ws_res.cell(row_idx, 9, f'=IF(H{row_idx}=0, "✅ PAS / KLOP", IF(H{row_idx}<0, "❌ KURANG (Tekor)", "⚠️ LEBIH"))')
        ws_res.cell(row_idx, 10, f'=IF(H{row_idx}=0, IF(F{row_idx}>0, "Setor sisa Rp " & TEXT(F{row_idx}, "#,##0") & " ke Kasir", "Lunas / Pas"), IF(H{row_idx}<0, "Wajib ganti tekor Rp " & TEXT(ABS(H{row_idx}), "#,##0"), "Setor kelebihan Rp " & TEXT(H{row_idx}, "#,##0")))')

        for c in range(1, 11):
            ws_res.cell(row_idx, c).border = thin_border
            if c in [1, 9]:
                ws_res.cell(row_idx, c).alignment = Alignment(horizontal='center', vertical='center')
            elif c in [4, 5, 6, 7, 8]:
                ws_res.cell(row_idx, c).alignment = Alignment(horizontal='right', vertical='center')
            else:
                ws_res.cell(row_idx, c).alignment = Alignment(horizontal='left', vertical='center')
        row_idx += 1

    # Total Row
    total_row = row_idx
    ws_res.cell(total_row, 1, "TOTAL KESELURUHAN TIM").font = bold_font
    ws_res.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=3)
    ws_res.cell(total_row, 1).alignment = Alignment(horizontal='center', vertical='center')

    ws_res.cell(total_row, 4, f"=SUM(D14:D{total_row-1})").number_format = rupiah_format
    ws_res.cell(total_row, 5, f"=SUM(E14:E{total_row-1})").number_format = rupiah_format
    ws_res.cell(total_row, 6, f"=SUM(F14:F{total_row-1})").number_format = rupiah_format
    ws_res.cell(total_row, 7, f"=SUM(G14:G{total_row-1})").number_format = rupiah_format
    ws_res.cell(total_row, 8, f"=SUM(H14:H{total_row-1})").number_format = rupiah_format
    ws_res.cell(total_row, 9, f'=IF(H{total_row}=0, "✅ FISIK TIM KLOP", IF(H{total_row}<0, "❌ TOTAL TEKOR " & TEXT(ABS(H{total_row}), "#,##0"), "⚠️ TOTAL LEBIH"))')
    ws_res.cell(total_row, 10, f'=IF(H{total_row}=0, "Semua fisik sesuai catatan", "Investigasi selisih fisik!")')

    for c in range(1, 11):
        cell = ws_res.cell(total_row, c)
        cell.fill = sub_header_fill
        cell.font = bold_font
        cell.border = thin_border
        if c in [4, 5, 6, 7, 8]:
            cell.alignment = Alignment(horizontal='right', vertical='center')
        elif c in [1, 9]:
            cell.alignment = Alignment(horizontal='center', vertical='center')

    # Audit Keseimbangan Kas
    audit_head_row = total_row + 2
    ws_res.cell(audit_head_row, 1, "KONTROL KESEIMBANGAN & AUDIT DANA PERSEKOT (BALANCE CHECK)").font = Font(name='Calibri', size=11, bold=True, color='1E3A8A')

    audit_items = [
        ("1. Plafon Kas Induk Persekot (Budget Awal)", f"=D7", "Plafon dana awal yang disediakan manajemen"),
        ("2. Total Realisasi Belanja Tim (Nota Sah)", f"=E{total_row}", "Total kuitansi/nota belanja tim yang sudah keluar"),
        ("3. Total Sisa Uang Fisik Riil Tim (Hasil Opname)", f"=G{total_row}", "Total lembaran/koin fisik uang di tangan seluruh anggota"),
        ("4. Sisa Fisik Dana di Kas Induk (Belum Dibagi)", f"=D9", "Fisik uang tunai yang masih tertahan di kasir utama"),
        ("5. Total Pertanggungjawaban Fisik (Belanja + Uang Fisik)", f"=E{audit_head_row+2}+E{audit_head_row+3}+E{audit_head_row+4}", "Total Uang Fisik + Nota (Wajib tepat sama dengan Plafon Kas)"),
        ("6. Selisih / Gap Keseimbangan Kas (Audit Variance)", f"=E{audit_head_row+5}-E{audit_head_row+1}", "Wajib bernilai Rp 0 agar dinyatakan Klop / Balance"),
    ]

    for offset, (label, formula_val, desc) in enumerate(audit_items, 1):
        cur_r = audit_head_row + offset
        ws_res.cell(cur_r, 2, label).font = bold_font
        c_val = ws_res.cell(cur_r, 5, formula_val)
        c_val.font = bold_font
        c_val.number_format = rupiah_format
        c_val.alignment = Alignment(horizontal='right')
        ws_res.cell(cur_r, 7, desc).font = subtitle_font
        ws_res.cell(cur_r, 2).fill = card_fill
        c_val.fill = card_fill

    status_audit_row = audit_head_row + 7
    ws_res.cell(status_audit_row, 2, "STATUS KESEIMBANGAN AUDIT KAS (BALANCE CHECK)").font = Font(name='Calibri', size=11, bold=True, color='1E3A8A')
    ws_res.cell(status_audit_row, 5, f'=IF(E{audit_head_row+6}=0, "✅ 100% BALANCE (Uang Fisik + Nota Lengkap Rp " & TEXT(E{audit_head_row+1}, "#,##0") & ")", IF(E{audit_head_row+6}<0, "❌ FISIK TEKOR / KURANG " & TEXT(ABS(E{audit_head_row+6}), "Rp #,##0"), "⚠️ FISIK LEBIH " & TEXT(E{audit_head_row+6}, "Rp #,##0")))').font = bold_font
    ws_res.cell(status_audit_row, 7, "Formula otomatis mendeteksi selisih fisik vs plafon").font = subtitle_font

    # --- MEMBER SHEETS ---
    for m in members:
        sheet_title = m["name"].replace("'", "")[:31]
        ws_m = wb.create_sheet(title=sheet_title)

        ws_m['A1'] = f"BUKU PENGELUARAN & REIMBURSEMENT PERSEKOT - {m['name'].upper()}"
        ws_m['A1'].font = title_font

        ws_m['A2'] = f"Nama Pemegang Kas: {m['name']}   |   Divisi: {m['divisi']}   |   Pencatatan Nota & Kas Bon"
        ws_m['A2'].font = subtitle_font

        # Get transactions for this member
        txs = get_transactions(member_id=m["id"])

        # In member sheet, we have summary KPI on rows 4-9
        # Row 4: Total Persekot Diterima -> =F{last_row}
        # Row 5: Total Realisasi Belanja (Nota) -> =G{last_row}
        # Row 6: Sisa Saldo Catatan (Buku) -> =D4-D5
        # Row 7: Sisa Uang Fisik Riil (Hasil Opname) -> =Resume!G{member_row_in_resume}
        # Row 8: Selisih / Gap Fisik (Fisik - Buku) -> =D7-D6
        # Row 9: Status Keseimbangan Fisik -> =IF(D8=0, "✅ PAS / KLOP (Uang Fisik Sesuai)", IF(D8<0, "❌ FISIK TEKOR / KURANG " & TEXT(ABS(D8), "Rp #,##0"), "⚠️ FISIK LEBIH"))

        member_resume_row = member_row_map.get(m["name"], 14)

        # We will write transactions starting from row 12
        start_tx_row = 12
        num_tx = max(len(txs), 1)
        end_tx_row = start_tx_row + num_tx - 1
        total_m_row = end_tx_row + 1

        ws_m['B4'] = "Total Persekot Diterima"
        ws_m['D4'] = f"=F{total_m_row}"
        ws_m['D4'].number_format = rupiah_format
        ws_m['D4'].font = bold_font

        ws_m['B5'] = "Total Realisasi Belanja (Nota)"
        ws_m['D5'] = f"=G{total_m_row}"
        ws_m['D5'].number_format = rupiah_format
        ws_m['D5'].font = bold_font

        ws_m['B6'] = "Sisa Saldo Catatan (Buku)"
        ws_m['D6'] = "=D4-D5"
        ws_m['D6'].number_format = rupiah_format
        ws_m['D6'].font = bold_font

        ws_m['B7'] = "Sisa Uang Fisik Riil (Hasil Opname)"
        ws_m['D7'] = f"=Resume!G{member_resume_row}"
        ws_m['D7'].number_format = rupiah_format
        ws_m['D7'].font = bold_font

        ws_m['B8'] = "Selisih / Gap Fisik (Fisik - Buku)"
        ws_m['D8'] = "=D7-D6"
        ws_m['D8'].number_format = rupiah_format
        ws_m['D8'].font = bold_font

        ws_m['B9'] = "Status Keseimbangan Fisik"
        ws_m['D9'] = '=IF(D8=0, "✅ PAS / KLOP (Uang Fisik Sesuai)", IF(D8<0, "❌ FISIK TEKOR / KURANG " & TEXT(ABS(D8), "Rp #,##0"), "⚠️ FISIK LEBIH"))'
        ws_m['D9'].font = bold_font

        for r in range(4, 10):
            ws_m.cell(r, 2).font = normal_font
            ws_m.cell(r, 2).fill = card_fill
            ws_m.cell(r, 4).fill = card_fill

        # Table Header
        headers_member = [
            "No.", "Tanggal", "No. Bukti / Nota", "Keterangan / Uraian Belanja",
            "Pos / Kategori Biaya", "Persekot Masuk (Rp)", "Pengeluaran / Nota (Rp)",
            "Saldo Berjalan (Rp)", "Status Bukti Fisik"
        ]

        for c, h in enumerate(headers_member, 1):
            cell = ws_m.cell(11, c, h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = thin_border
        ws_m.row_dimensions[11].height = 25

        # Transaction Rows
        for i, t in enumerate(txs, 1):
            r = start_tx_row + i - 1
            ws_m.cell(r, 1, i).alignment = Alignment(horizontal='center')
            ws_m.cell(r, 2, t["tanggal"]).alignment = Alignment(horizontal='center')
            ws_m.cell(r, 3, t["no_bukti"]).alignment = Alignment(horizontal='center')
            ws_m.cell(r, 4, t["keterangan"])
            ws_m.cell(r, 5, t["kategori"])

            c_in = ws_m.cell(r, 6, t["masuk"] if t["masuk"] > 0 else None)
            c_in.number_format = rupiah_format
            c_in.alignment = Alignment(horizontal='right')

            c_out = ws_m.cell(r, 7, t["keluar"] if t["keluar"] > 0 else None)
            c_out.number_format = rupiah_format
            c_out.alignment = Alignment(horizontal='right')

            # Formula for running balance: =IF(AND(F12="", G12=""), "", SUM($F$12:F12)-SUM($G$12:G12))
            c_sal = ws_m.cell(r, 8, f'=IF(AND(F{r}="", G{r}=""), "", SUM($F$12:F{r})-SUM($G$12:G{r}))')
            c_sal.number_format = rupiah_format
            c_sal.alignment = Alignment(horizontal='right')

            ws_m.cell(r, 9, t["status_bukti"]).alignment = Alignment(horizontal='center')

            for c in range(1, 10):
                ws_m.cell(r, c).border = thin_border
                ws_m.cell(r, c).font = normal_font

        # Total Row for Member
        ws_m.cell(total_m_row, 1, "TOTAL PERSEKOT & REALISASI PENGELUARAN").font = bold_font
        ws_m.merge_cells(start_row=total_m_row, start_column=1, end_row=total_m_row, end_column=5)
        ws_m.cell(total_m_row, 1).alignment = Alignment(horizontal='center', vertical='center')

        ws_m.cell(total_m_row, 6, f"=SUM(F12:F{end_tx_row})").number_format = rupiah_format
        ws_m.cell(total_m_row, 7, f"=SUM(G12:G{end_tx_row})").number_format = rupiah_format
        ws_m.cell(total_m_row, 8, f"=F{total_m_row}-G{total_m_row}").number_format = rupiah_format

        for c in range(1, 10):
            cell = ws_m.cell(total_m_row, c)
            cell.fill = sub_header_fill
            cell.font = bold_font
            cell.border = thin_border
            if c in [6, 7, 8]:
                cell.alignment = Alignment(horizontal='right', vertical='center')

        # Auto-fit column widths
        for col in ws_m.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_m.column_dimensions[col_letter].width = max(max_len + 3, 12)
        ws_m.column_dimensions['D'].width = 32

    # Auto-fit Resume columns
    for col in ws_res.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_res.column_dimensions[col_letter].width = max(max_len + 3, 12)
    ws_res.column_dimensions['B'].width = 38
    ws_res.column_dimensions['D'].width = 24
    ws_res.column_dimensions['J'].width = 34

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def import_excel_to_database(file_bytes_or_path):
    wb = openpyxl.load_workbook(file_bytes_or_path, data_only=True)
    conn = get_connection()
    cur = conn.cursor()

    # Reset
    cur.execute("DELETE FROM transactions")
    cur.execute("DELETE FROM members")

    ws_res = wb['Resume'] if 'Resume' in wb.sheetnames else wb.active

    # Plafon
    plafon_val = ws_res.cell(7, 4).value
    if plafon_val:
        try:
            cur.execute("INSERT OR REPLACE INTO settings (key, value) VALUES ('plafon_induk', ?)", (str(int(float(plafon_val))),))
        except Exception:
            pass

    members_map = {}
    for r in range(14, ws_res.max_row + 1):
        name = ws_res.cell(r, 2).value
        if not name or 'TOTAL' in str(name).upper():
            break
        divisi = ws_res.cell(r, 3).value or ''
        fisik = float(ws_res.cell(r, 7).value or 0)
        cur.execute("INSERT INTO members (name, divisi, uang_fisik, urutan) VALUES (?, ?, ?, ?)",
                    (str(name).strip(), str(divisi).strip(), fisik, r - 13))
        members_map[str(name).strip().lower()] = cur.lastrowid

    # Transactions from each member sheet
    for name_lower, m_id in members_map.items():
        sheet_name = None
        for s in wb.sheetnames:
            if s.strip().lower() == name_lower:
                sheet_name = s
                break
        if not sheet_name:
            continue
        ws = wb[sheet_name]
        for r in range(12, ws.max_row + 1):
            no_val = ws.cell(r, 1).value
            if isinstance(no_val, str) and 'TOTAL' in no_val.upper():
                break
            tgl = ws.cell(r, 2).value
            bukti = ws.cell(r, 3).value
            ket = ws.cell(r, 4).value
            pos = ws.cell(r, 5).value
            masuk = float(ws.cell(r, 6).value or 0)
            keluar = float(ws.cell(r, 7).value or 0)
            status_bukti = ws.cell(r, 9).value or 'Kas Diterima'
            if tgl or bukti or ket or masuk > 0 or keluar > 0:
                tgl_str = str(tgl)[:10] if tgl else '2026-01-09'
                cur.execute("""
                    INSERT INTO transactions (member_id, tanggal, no_bukti, keterangan, kategori, masuk, keluar, status_bukti)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """, (m_id, tgl_str, str(bukti or ''), str(ket or ''), str(pos or 'Operasional'), masuk, keluar, str(status_bukti)))

    conn.commit()
    conn.close()
    return True
