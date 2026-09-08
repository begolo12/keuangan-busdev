import sqlite3
import os
import openpyxl
from datetime import datetime

DB_PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "keuangan_busdev.db")
EXCEL_FILE = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "Keuangan Busdev (1).xlsx")

def get_connection():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn

def init_db():
    conn = get_connection()
    cur = conn.cursor()

    cur.execute("""
    CREATE TABLE IF NOT EXISTS settings (
        key TEXT PRIMARY KEY,
        value TEXT
    );
    """)

    cur.execute("""
    CREATE TABLE IF NOT EXISTS members (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT NOT NULL,
        divisi TEXT NOT NULL,
        uang_fisik REAL DEFAULT 0,
        urutan INTEGER DEFAULT 0,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    """)

    cur.execute("""
    CREATE TABLE IF NOT EXISTS transactions (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        member_id INTEGER NOT NULL,
        tanggal TEXT NOT NULL,
        no_bukti TEXT NOT NULL,
        keterangan TEXT NOT NULL,
        kategori TEXT NOT NULL,
        masuk REAL DEFAULT 0,
        keluar REAL DEFAULT 0,
        status_bukti TEXT DEFAULT 'Nota Lengkap',
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY (member_id) REFERENCES members (id) ON DELETE CASCADE
    );
    """)

    conn.commit()

    # Check if empty, then seed from Excel or default
    cur.execute("SELECT COUNT(*) FROM members")
    count = cur.fetchone()[0]
    if count == 0:
        seed_from_excel(conn)

    conn.close()

def seed_from_excel(conn=None):
    close_at_end = False
    if conn is None:
        conn = get_connection()
        close_at_end = True

    cur = conn.cursor()
    # Reset existing
    cur.execute("DELETE FROM transactions")
    cur.execute("DELETE FROM members")
    cur.execute("DELETE FROM settings")

    # Default settings
    cur.execute("INSERT OR REPLACE INTO settings (key, value) VALUES ('plafon_induk', '7000000')")
    cur.execute("INSERT OR REPLACE INTO settings (key, value) VALUES ('tahun', '2026')")
    cur.execute("INSERT OR REPLACE INTO settings (key, value) VALUES ('judul', 'DASHBOARD REKAPITULASI DANA PERSEKOT, REIMBURSEMENT & OPNAME KAS FISIK')")

    if os.path.exists(EXCEL_FILE):
        try:
            wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
            ws_res = wb['Resume']

            # Extract plafon from Resume if exists
            plafon_val = ws_res.cell(7, 4).value
            if plafon_val:
                try:
                    cur.execute("INSERT OR REPLACE INTO settings (key, value) VALUES ('plafon_induk', ?)", (str(int(float(plafon_val))),))
                except Exception:
                    pass

            members_map = {}
            for r in range(14, 18):
                name = ws_res.cell(r, 2).value
                divisi = ws_res.cell(r, 3).value
                fisik = float(ws_res.cell(r, 7).value or 0)
                if name:
                    cur.execute("INSERT INTO members (name, divisi, uang_fisik, urutan) VALUES (?, ?, ?, ?)",
                                (str(name).strip(), str(divisi or '').strip(), fisik, r - 13))
                    members_map[str(name).strip().lower()] = cur.lastrowid

            for name_lower, m_id in members_map.items():
                sheet_name = None
                for s in wb.sheetnames:
                    if s.strip().lower() == name_lower:
                        sheet_name = s
                        break
                if not sheet_name:
                    continue
                ws = wb[sheet_name]
                for r in range(12, ws.max_row + 1):
                    no_val = ws.cell(r, 1).value
                    if isinstance(no_val, str) and 'TOTAL' in no_val.upper():
                        break
                    tgl = ws.cell(r, 2).value
                    bukti = ws.cell(r, 3).value
                    ket = ws.cell(r, 4).value
                    pos = ws.cell(r, 5).value
                    masuk = float(ws.cell(r, 6).value or 0)
                    keluar = float(ws.cell(r, 7).value or 0)
                    status_bukti = ws.cell(r, 9).value or 'Kas Diterima'
                    if tgl or bukti or ket or masuk > 0 or keluar > 0:
                        tgl_str = str(tgl)[:10] if tgl else '2026-01-09'
                        cur.execute("""
                            INSERT INTO transactions (member_id, tanggal, no_bukti, keterangan, kategori, masuk, keluar, status_bukti)
                            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                        """, (m_id, tgl_str, str(bukti or ''), str(ket or ''), str(pos or 'Operasional'), masuk, keluar, str(status_bukti)))
            conn.commit()
        except Exception as e:
            print("Error loading excel:", e)
            seed_defaults(conn)
    else:
        seed_defaults(conn)

    if close_at_end:
        conn.close()

def seed_defaults(conn):
    cur = conn.cursor()
    members_data = [
        ('Sri', 'Dapur', 2500000, 1, 2500000),
        ('Lucky', 'GA', 3200000, 2, 3200000),
        ('Kur', 'OB', 300000, 3, 300000),
        ('irvan', 'Manager', 1000000, 4, 1000000),
    ]
    for name, divisi, fisik, urutan, persekot in members_data:
        cur.execute("INSERT INTO members (name, divisi, uang_fisik, urutan) VALUES (?, ?, ?, ?)",
                    (name, divisi, fisik, urutan))
        m_id = cur.lastrowid
        cur.execute("""
            INSERT INTO transactions (member_id, tanggal, no_bukti, keterangan, kategori, masuk, keluar, status_bukti)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """, (m_id, '2026-01-09', 'KAS-01', 'Penerimaan Persekot Awal', 'Persekot Kas', persekot, 0, 'Kas Diterima'))
    conn.commit()

def get_settings():
    conn = get_connection()
    rows = conn.execute("SELECT key, value FROM settings").fetchall()
    conn.close()
    return {row["key"]: row["value"] for row in rows}

def update_setting(key, value):
    conn = get_connection()
    conn.execute("INSERT OR REPLACE INTO settings (key, value) VALUES (?, ?)", (key, str(value)))
    conn.commit()
    conn.close()

def get_members():
    conn = get_connection()
    members = conn.execute("SELECT * FROM members ORDER BY urutan ASC, id ASC").fetchall()
    result = []
    for m in members:
        # Calculate stats
        stats = conn.execute("""
            SELECT 
                COALESCE(SUM(masuk), 0) as total_masuk,
                COALESCE(SUM(keluar), 0) as total_keluar
            FROM transactions WHERE member_id = ?
        """, (m["id"],)).fetchone()

        total_masuk = float(stats["total_masuk"])
        total_keluar = float(stats["total_keluar"])
        saldo_buku = total_masuk - total_keluar
        uang_fisik = float(m["uang_fisik"])
        selisih = uang_fisik - saldo_buku

        if selisih == 0:
            status_fisik = "✅ SEIMBANG"
            status_code = "seimbang"
            if saldo_buku > 0:
                tindak_lanjut = f"Setor sisa Rp {saldo_buku:,.0f} ke Kasir".replace(",", ".")
            else:
                tindak_lanjut = "Sudah Tertentuk"
        elif selisih < 0:
            status_fisik = "❌ KEKURANGAN"
            status_code = "kekurangan"
            tindak_lanjut = f"Wajib mengganti kekurangan Rp {abs(selisih):,.0f}".replace(",", ".")
        else:
            status_fisik = "⚠️ BERLEBIHAN"
            status_code = "berlebihan"
            tindak_lanjut = f"Setor kelebihan Rp {selisih:,.0f}".replace(",", ".")

        result.append({
            "id": m["id"],
            "name": m["name"],
            "divisi": m["divisi"],
            "uang_fisik": uang_fisik,
            "urutan": m["urutan"],
            "total_masuk": total_masuk,
            "total_keluar": total_keluar,
            "saldo_buku": saldo_buku,
            "selisih": selisih,
            "status_fisik": status_fisik,
            "status_code": status_code,
            "tindak_lanjut": tindak_lanjut
        })
    conn.close()
    return result

def get_member(member_id):
    conn = get_connection()
    m = conn.execute("SELECT * FROM members WHERE id = ?", (member_id,)).fetchone()
    if not m:
        conn.close()
        return None
    stats = conn.execute("""
        SELECT 
            COALESCE(SUM(masuk), 0) as total_masuk,
            COALESCE(SUM(keluar), 0) as total_keluar
        FROM transactions WHERE member_id = ?
    """, (m["id"],)).fetchone()
    conn.close()

    total_masuk = float(stats["total_masuk"])
    total_keluar = float(stats["total_keluar"])
    saldo_buku = total_masuk - total_keluar
    uang_fisik = float(m["uang_fisik"])
    selisih = uang_fisik - saldo_buku

    if selisih == 0:
        status_fisik = "✅ SEIMBANG"
        status_code = "seimbang"
        tindak_lanjut = f"Setor sisa Rp {saldo_buku:,.0f} ke Kasir".replace(",", ".") if saldo_buku > 0 else "Sudah Tertentuk"
    elif selisih < 0:
        status_fisik = "❌ KEKURANGAN"
        status_code = "kekurangan"
        tindak_lanjut = f"Wajib mengganti kekurangan Rp {abs(selisih):,.0f}".replace(",", ".")
    else:
        status_fisik = "⚠️ BERLEBIHAN"
        status_code = "berlebihan"
        tindak_lanjut = f"Setor kelebihan Rp {selisih:,.0f}".replace(",", ".")
    return {
        "id": m["id"],
        "name": m["name"],
        "divisi": m["divisi"],
        "uang_fisik": uang_fisik,
        "urutan": m["urutan"],
        "total_masuk": total_masuk,
        "total_keluar": total_keluar,
        "saldo_buku": saldo_buku,
        "selisih": selisih,
        "status_fisik": status_fisik,
        "status_code": status_code,
        "tindak_lanjut": tindak_lanjut
    }

def add_member(name, divisi, uang_fisik=0, urutan=0):
    conn = get_connection()
    cur = conn.cursor()
    cur.execute("INSERT INTO members (name, divisi, uang_fisik, urutan) VALUES (?, ?, ?, ?)",
                (name.strip(), divisi.strip(), float(uang_fisik or 0), int(urutan or 0)))
    conn.commit()
    new_id = cur.lastrowid
    conn.close()
    return new_id

def update_member(member_id, name, divisi, uang_fisik=None, urutan=None):
    conn = get_connection()
    cur = conn.cursor()
    if uang_fisik is not None and urutan is not None:
        cur.execute("UPDATE members SET name = ?, divisi = ?, uang_fisik = ?, urutan = ? WHERE id = ?",
                    (name.strip(), divisi.strip(), float(uang_fisik), int(urutan), member_id))
    elif uang_fisik is not None:
        cur.execute("UPDATE members SET name = ?, divisi = ?, uang_fisik = ? WHERE id = ?",
                    (name.strip(), divisi.strip(), float(uang_fisik), member_id))
    else:
        cur.execute("UPDATE members SET name = ?, divisi = ? WHERE id = ?",
                    (name.strip(), divisi.strip(), member_id))
    conn.commit()
    conn.close()

def delete_member(member_id):
    conn = get_connection()
    conn.execute("DELETE FROM members WHERE id = ?", (member_id,))
    conn.commit()
    conn.close()

def update_member_opname(member_id, uang_fisik):
    conn = get_connection()
    conn.execute("UPDATE members SET uang_fisik = ? WHERE id = ?", (float(uang_fisik or 0), member_id))
    conn.commit()
    conn.close()

def get_transactions(member_id=None, search=None, kategori=None, jenis=None, status_bukti=None, date_from=None, date_to=None):
    conn = get_connection()
    query = """
        SELECT t.*, m.name as member_name, m.divisi as member_divisi
        FROM transactions t
        JOIN members m ON t.member_id = m.id
        WHERE 1=1
    """
    params = []

    if member_id:
        query += " AND t.member_id = ?"
        params.append(member_id)
    if search:
        query += " AND (t.keterangan LIKE ? OR t.no_bukti LIKE ?)"
        params.extend([f"%{search}%", f"%{search}%"])
    if kategori:
        query += " AND t.kategori = ?"
        params.append(kategori)
    if jenis == "masuk":
        query += " AND t.masuk > 0"
    elif jenis == "keluar":
        query += " AND t.keluar > 0"
    if status_bukti:
        query += " AND t.status_bukti = ?"
        params.append(status_bukti)
    if date_from:
        query += " AND t.tanggal >= ?"
        params.append(date_from)
    if date_to:
        query += " AND t.tanggal <= ?"
        params.append(date_to)

    query += " ORDER BY t.tanggal ASC, t.id ASC"
    rows = conn.execute(query, params).fetchall()

    result = []
    if member_id:
        running_balance = 0
        for r in rows:
            masuk = float(r["masuk"] or 0)
            keluar = float(r["keluar"] or 0)
            running_balance += (masuk - keluar)
            result.append({
                "id": r["id"],
                "member_id": r["member_id"],
                "member_name": r["member_name"],
                "member_divisi": r["member_divisi"],
                "tanggal": r["tanggal"],
                "no_bukti": r["no_bukti"],
                "keterangan": r["keterangan"],
                "kategori": r["kategori"],
                "masuk": masuk,
                "keluar": keluar,
                "saldo_berjalan": running_balance,
                "status_bukti": r["status_bukti"],
                "created_at": r["created_at"]
            })
    else:
        member_balances = {}
        for r in rows:
            m_id = r["member_id"]
            masuk = float(r["masuk"] or 0)
            keluar = float(r["keluar"] or 0)
            member_balances[m_id] = member_balances.get(m_id, 0) + (masuk - keluar)
            result.append({
                "id": r["id"],
                "member_id": r["member_id"],
                "member_name": r["member_name"],
                "member_divisi": r["member_divisi"],
                "tanggal": r["tanggal"],
                "no_bukti": r["no_bukti"],
                "keterangan": r["keterangan"],
                "kategori": r["kategori"],
                "masuk": masuk,
                "keluar": keluar,
                "saldo_berjalan": member_balances[m_id],
                "status_bukti": r["status_bukti"],
                "created_at": r["created_at"]
            })
    conn.close()
    return result

def get_transaction(tx_id):
    conn = get_connection()
    row = conn.execute("""
        SELECT t.*, m.name as member_name, m.divisi as member_divisi
        FROM transactions t
        JOIN members m ON t.member_id = m.id
        WHERE t.id = ?
    """, (tx_id,)).fetchone()
    conn.close()
    if not row:
        return None
    return dict(row)

def add_transaction(member_id, tanggal, no_bukti, keterangan, kategori, masuk, keluar, status_bukti):
    conn = get_connection()
    cur = conn.cursor()
    cur.execute("""
        INSERT INTO transactions (member_id, tanggal, no_bukti, keterangan, kategori, masuk, keluar, status_bukti)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
    """, (member_id, tanggal.strip(), no_bukti.strip(), keterangan.strip(), kategori.strip(),
          float(masuk or 0), float(keluar or 0), status_bukti.strip()))
    conn.commit()
    new_id = cur.lastrowid
    conn.close()
    return new_id

def update_transaction(tx_id, member_id, tanggal, no_bukti, keterangan, kategori, masuk, keluar, status_bukti):
    conn = get_connection()
    cur = conn.cursor()
    cur.execute("""
        UPDATE transactions
        SET member_id = ?, tanggal = ?, no_bukti = ?, keterangan = ?, kategori = ?, masuk = ?, keluar = ?, status_bukti = ?
        WHERE id = ?
    """, (member_id, tanggal.strip(), no_bukti.strip(), keterangan.strip(), kategori.strip(),
          float(masuk or 0), float(keluar or 0), status_bukti.strip(), tx_id))
    conn.commit()
    conn.close()

def delete_transaction(tx_id):
    conn = get_connection()
    conn.execute("DELETE FROM transactions WHERE id = ?", (tx_id,))
    conn.commit()
    conn.close()

def get_categories():
    conn = get_connection()
    rows = conn.execute("SELECT DISTINCT kategori FROM transactions WHERE kategori IS NOT NULL AND kategori != '' ORDER BY kategori ASC").fetchall()
    conn.close()
    categories = [r[0] for r in rows]
    # Default set if minimal
    defaults = ["Persekot Kas", "Dapur / Konsumsi", "Operasional GA", "Kebersihan / Alat OB", "Transport / Bensin", "Operasional Bisnis", "Lain-lain"]
    for d in defaults:
        if d not in categories:
            categories.append(d)
    return categories

def get_summary():
    settings = get_settings()
    plafon_induk = float(settings.get("plafon_induk", 7000000))
    members = get_members()

    total_distribusi = sum(m["total_masuk"] for m in members)
    sisa_kas_induk = plafon_induk - total_distribusi

    if sisa_kas_induk == 0:
        status_distribusi = "100% Teralokasi ke Tim"
        distribusi_code = "full"
    elif sisa_kas_induk > 0:
        status_distribusi = "Sebagian Dana Masih di Kas Induk"
        distribusi_code = "partial"
    else:
        status_distribusi = "PERINGATAN: Melebihi Plafon Kas!"
        distribusi_code = "over"

    total_realisasi_belanja = sum(m["total_keluar"] for m in members)
    total_saldo_buku = sum(m["saldo_buku"] for m in members)
    total_uang_fisik = sum(m["uang_fisik"] for m in members)
    total_selisih_fisik = sum(m["selisih"] for m in members)

    if total_selisih_fisik == 0:
        status_selisih_tim = "✅ SELURUH ANGGOTA SEIMBANG"
        tindak_lanjut_tim = "Semua fisik sesuai catatan"
    elif total_selisih_fisik < 0:
        status_selisih_tim = f"❌ TOTAL KEKURANGAN Rp {abs(total_selisih_fisik):,.0f}".replace(",", ".")
        tindak_lanjut_tim = "Investigasi kekurangan fisik!"
    else:
        status_selisih_tim = f"⚠️ TOTAL BERLEBIHAN Rp {total_selisih_fisik:,.0f}".replace(",", ".")
        tindak_lanjut_tim = "Investigasi kelebihan fisik!"
    # Audit Balance Check
    # Total Pertanggungjawaban Fisik = Realisasi Belanja + Sisa Uang Fisik Tim + Sisa Kas Induk
    audit_total = total_realisasi_belanja + total_uang_fisik + sisa_kas_induk
    audit_gap = audit_total - plafon_induk

    if audit_gap == 0:
        status_audit = f"✅ 100% BALANCE (Uang Fisik + Nota Lengkap Rp {plafon_induk:,.0f})".replace(",", ".")
        audit_code = "balance"
    elif audit_gap < 0:
        status_audit = f"❌ KEKURANGAN FISIK Rp {abs(audit_gap):,.0f}".replace(",", ".")
        audit_code = "kekurangan"
    else:
        status_audit = f"⚠️ BERLEBIHAN FISIK Rp {audit_gap:,.0f}".replace(",", ".")
    return {
        "plafon_induk": plafon_induk,
        "total_distribusi": total_distribusi,
        "sisa_kas_induk": sisa_kas_induk,
        "status_distribusi": status_distribusi,
        "distribusi_code": distribusi_code,
        "total_realisasi_belanja": total_realisasi_belanja,
        "total_saldo_buku": total_saldo_buku,
        "total_uang_fisik": total_uang_fisik,
        "total_selisih_fisik": total_selisih_fisik,
        "status_selisih_tim": status_selisih_tim,
        "tindak_lanjut_tim": tindak_lanjut_tim,
        "audit_total_pertanggungjawaban": audit_total,
        "audit_gap": audit_gap,
        "status_audit": status_audit,
        "audit_code": audit_code,
        "members": members,
        "tahun": settings.get("tahun", "2026"),
        "judul": settings.get("judul", "DASHBOARD REKAPITULASI DANA PERSEKOT, REIMBURSEMENT & OPNAME KAS FISIK")
    }
